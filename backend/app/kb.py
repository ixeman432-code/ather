"""
Athar Hadith KB
Source: approved project Excel workbook.
The workbook is read once when this module starts.
"""

from pathlib import Path
import openpyxl

EXCEL_PATH = (
    Path(__file__).resolve().parents[1]
    / "data"
    / "raw"
    / "hadith_data.xlsx"
)

RESULT_MAP = {
    "excellent": "ثناء",
    "neutral": "تشجيع",
    "poor": "تنبيه لطيف",
}

HADITHS = []


def _load_hadiths():
    if not EXCEL_PATH.exists():
        return []

    wb = openpyxl.load_workbook(
        EXCEL_PATH,
        read_only=True,
        data_only=True,
    )

    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    headers = next(rows, None)

    if not headers:
        wb.close()
        return []

    headers = [
        str(value).strip() if value is not None else ""
        for value in headers
    ]

    data = []

    for values in rows:
        row = dict(zip(headers, values))

        if str(row.get("معتمد؟") or "").strip() != "نعم":
            continue

        text = str(row.get("نص العرض") or "").strip()
        source = str(row.get("رابط المصدر") or "").strip()

        if not text or not source:
            continue

        target_age = row.get("العمر المستهدف")

        try:
            target_age = int(float(target_age))
        except (TypeError, ValueError):
            continue

        data.append({
            "id": str(row.get("رقم الحديث") or "").strip(),
            "value": str(row.get("القيمة") or "").strip(),
            "text": text,
            "source": source,
            "text_en": str(row.get("الترجمة الإنجليزية") or "").strip(),
            "tone": str(row.get("النبرة") or "").strip(),
            "age_target": target_age,
            "approved_en": (
                str(row.get("ترجمة معتمدة") or "").strip() == "نعم"
                if "ترجمة معتمدة" in headers
                else False
            ),
        })

    wb.close()
    return data


HADITHS = _load_hadiths()
print(f"ATHAR DEBUG: loaded hadiths = {len(HADITHS)}")
print(f"ATHAR DEBUG: Excel exists = {EXCEL_PATH.exists()}")


def get_hadith_pool(classification, child_age, language="ar"):
    result = RESULT_MAP.get(classification)

    if not result:
        return []

    try:
        age = int(child_age)
    except (TypeError, ValueError):
        return []

    if age <= 9:
        target_age = 9
    elif age <= 12:
        target_age = 10
    else:
        return []

    if language == "en":
        return [
            hadith
            for hadith in HADITHS
            if hadith["tone"] == result
            and hadith["age_target"] == target_age
            and hadith["approved_en"]
            and hadith["text_en"]
        ]

    return [
        hadith
        for hadith in HADITHS
        if hadith["tone"] == result
        and hadith["age_target"] == target_age
    ]
